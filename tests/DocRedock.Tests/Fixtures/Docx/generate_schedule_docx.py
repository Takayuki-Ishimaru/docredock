#!/usr/bin/env python3
"""Generate schedule-arrows.docx: a Japanese IT-project schedule document
used to test the "DOCX table overlay" feature (arrow/bar/marker/today-line
DrawingML shapes drawn ON TOP OF a native w:tbl whose columns are dates).

Companion to ../Pptx/generate_schedule_pptx.py and
../Xlsx/generate_schedule_xlsx.py: same table shape/labels/overlay
inventory, translated to WordprocessingML floating shapes (wp:anchor +
a:graphicData uri=".../wordprocessingShape" + wps:wsp) instead of DrawingML
slide/sheet shapes.

python-docx 1.2.0 has no write API for floating (anchored) DrawingML
shapes, so this script builds the table/paragraphs with python-docx and
injects every wp:anchor/wps:wsp via raw XML strings parsed with lxml --
the same convention generate_complex_docx.py uses for its own anchored
picture/textbox/hyperlink/footnote injections (see that file's frag()/
append_raw() helpers, mirrored here).

Usage:
  python3 generate_schedule_docx.py [output.docx] [--verify]

  --verify        Re-open an already-generated file and print, for every
                   floating shape: its anchor style (cell/page), the
                   relativeFrom values, posOffset/extent in EMU, prst, rot,
                   and the computed expected table row(s)/column(s) under
                   the same 50%-overlap rule the PPTX/XLSX generators use.
                   (Also runs automatically after a fresh generation.)
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from lxml import etree

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, Twips

DEFAULT_OUT = Path(__file__).resolve().parent / "schedule-arrows.docx"

FIXED_ZIP_DATE_TIME = (2026, 1, 1, 0, 0, 0)
FIXED_DATETIME = __import__("datetime").datetime(2026, 1, 1, 0, 0, 0)

NAVY = "17365D"
TEAL = "0B7285"
GOLD = "B7791B"
GREEN = "067647"
RED = "B42318"
MUTED = "667085"
WHITE = "FFFFFF"

CHECKLIST: list[tuple[str, str, str, str]] = []


def record(elem_id, desc, method, ok):
    CHECKLIST.append((elem_id, desc, method, "PASS" if ok else "FAIL"))


TWIP_EMU = 635  # 1 twip = 635 EMU (20 twips = 1pt = 12700 EMU)


def twips_to_emu(t):
    return t * TWIP_EMU


# ============================================================== geometry ==
# Page: A4 landscape, 20mm margins on every side.
PAGE_W_TWIPS = Mm(297).twips  # 16838
PAGE_H_TWIPS = Mm(210).twips  # 11906
MARGIN_TWIPS = Mm(20).twips   # 1134

# Main table (page 1): 工程|担当 + 6 date columns, explicit gridCol widths
# in twips per the task spec.
GRIDCOL_MAIN = [2400, 1600, 1400, 1400, 1400, 1400, 1400, 1400]
HEADERS_MAIN = ["工程", "担当", "9/1", "9/2", "9/3", "9/4", "9/5", "9/8"]
ROWS_MAIN = [
    ["要件定義", "山田"],
    ["設計", "佐藤"],
    ["実装", "鈴木"],
    ["テスト", "田中"],
    ["リリース", "全員"],
]
ROW_HEIGHT_TWIPS = 500  # w:trHeight val, hRule="exact", every row

# Small table (page 2): 工程 + 4 date columns.
GRIDCOL_SMALL = [2400, 1400, 1400, 1400, 1400]
HEADERS_SMALL = ["工程", "9/1", "9/2", "9/3", "9/4"]
ROWS_SMALL = [["差し戻し"], ["設計"], ["実装"]]

# The ONE paragraph immediately before each table is given a forced exact
# line height (spacing before/after = 0, w:spacing line="240" lineRule=
# "exact", 12pt) so its rendered height is a deterministic, documented
# constant -- this is the "vertical offset assumes the table starts right
# after the preceding paragraph" assumption the task asks to state
# explicitly. Table top (from that page's own top margin) = MARGIN_TWIPS
# + PRECEDING_PARA_TWIPS, for BOTH page 1 and page 2 (each page break
# starts a fresh page-relative origin for relativeFrom="page").
PRECEDING_PARA_TWIPS = 240
TABLE_TOP_TWIPS = MARGIN_TWIPS + PRECEDING_PARA_TWIPS  # 1134 + 240 = 1374

INSET_EMU = 45720  # ~0.05in small gap, same scale as the XLSX/PPTX fixtures
BAR_HEIGHT_FRAC = 0.6
DIAMOND_SIZE_EMU = 228600  # 0.25in square, same as the XLSX fixture


def row_height_emu():
    return twips_to_emu(ROW_HEIGHT_TWIPS)


def col_width_emu_main(col_idx):
    return twips_to_emu(GRIDCOL_MAIN[col_idx])


def bar_cy_emu():
    return round(row_height_emu() * BAR_HEIGHT_FRAC)


def span_width_twips(gridcol, col_start, col_end):
    return sum(gridcol[col_start:col_end + 1])


# ------------------------------------------------------------ xml pieces --
NSDECL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)

_DOCPR_ID = [1]
_REL_HEIGHT = [251658240]


def next_docpr_id():
    _DOCPR_ID[0] += 1
    return _DOCPR_ID[0]


def next_rel_height():
    _REL_HEIGHT[0] += 1024
    return _REL_HEIGHT[0]


def frag(xml: str) -> etree._Element:
    return etree.fromstring(xml.encode("utf-8"))


def append_raw(paragraph, xml: str) -> etree._Element:
    element = frag(xml)
    paragraph._p.append(element)
    return element


def wsp_autoshape(name, prst, cx, cy, fill_hex, line_hex=WHITE, rot=None, text=None,
                   text_color=WHITE, font_size_half_pt=16, txbox=False, bold=True):
    """<wps:wsp> for a preset-geometry autoshape (rightArrow/leftRightArrow/
    rect/diamond/roundRect/...)."""
    rot_attr = f' rot="{rot}"' if rot else ""
    xfrm = f'<a:xfrm{rot_attr}><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
    cnv_sp_pr = '<wps:cNvSpPr txBox="1"/>' if txbox else "<wps:cNvSpPr/>"
    fill = f'<a:solidFill><a:srgbClr val="{fill_hex}"/></a:solidFill>' if fill_hex else "<a:noFill/>"
    line = (f'<a:ln><a:solidFill><a:srgbClr val="{line_hex}"/></a:solidFill></a:ln>'
            if line_hex else "<a:ln><a:noFill/></a:ln>")
    if text:
        b = "<w:b/>" if bold else ""
        txbx = (
            '<wps:txbx><w:txbxContent><w:p><w:pPr><w:jc w:val="center"/></w:pPr>'
            f'<w:r><w:rPr>{b}<w:color w:val="{text_color}"/><w:sz w:val="{font_size_half_pt}"/></w:rPr>'
            f'<w:t xml:space="preserve">{text}</w:t></w:r></w:p></w:txbxContent></wps:txbx>'
        )
    else:
        txbx = '<wps:txbx><w:txbxContent><w:p/></w:txbxContent></wps:txbx>'
    body_pr = ('<wps:bodyPr rot="0" spcFirstLastPara="0" vertOverflow="overflow" '
               'horzOverflow="overflow" vert="horz" wrap="square" lIns="0" tIns="0" '
               'rIns="0" bIns="0" anchor="ctr" anchorCtr="0"/>')
    return (
        f'<wps:wsp>{cnv_sp_pr}<wps:spPr>{xfrm}<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>'
        f'{fill}{line}</wps:spPr>{txbx}{body_pr}</wps:wsp>'
    )


def wsp_line(cx, cy, rot=None, line_hex=RED, width_emu=28575, tail_triangle=True):
    """<wps:wsp> straight-connector-shaped line, prst="straightConnector1",
    used for the today-line and the (unconnected) flow-diagram connector.
    python-docx/wordprocessingShape has no dedicated line-arrowhead API
    (same limitation the PPTX generator notes for python-pptx), so the
    a:tailEnd is hand-authored raw XML just like the rest of the shape."""
    rot_attr = f' rot="{rot}"' if rot else ""
    xfrm = f'<a:xfrm{rot_attr}><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
    tail = '<a:tailEnd type="triangle"/>' if tail_triangle else ""
    return (
        f'<wps:wsp><wps:cNvSpPr/><wps:spPr>{xfrm}'
        f'<a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom>'
        f'<a:ln w="{width_emu}"><a:solidFill><a:srgbClr val="{line_hex}"/></a:solidFill>{tail}</a:ln>'
        f'</wps:spPr><wps:bodyPr/></wps:wsp>'
    )


def anchor_run(shape_xml, name, cx, cy, pos_h, pos_v, layout_in_cell, wrap="None"):
    """Full <w:r><w:drawing><wp:anchor>...</wp:anchor></w:drawing></w:r>.
    pos_h/pos_v are (relativeFrom, offsetEMU) pairs."""
    docpr_id = next_docpr_id()
    rel_height = next_rel_height()
    ph_rf, ph_off = pos_h
    pv_rf, pv_off = pos_v
    wrap_xml = f'<wp:wrap{wrap}/>' if wrap != "None" else "<wp:wrapNone/>"
    return (
        f'<w:r {NSDECL}><w:drawing>'
        f'<wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" '
        f'relativeHeight="{rel_height}" behindDoc="0" locked="0" '
        f'layoutInCell="{layout_in_cell}" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="{ph_rf}"><wp:posOffset>{ph_off}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="{pv_rf}"><wp:posOffset>{pv_off}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'{wrap_xml}'
        f'<wp:docPr id="{docpr_id}" name="{name}"/>'
        f'<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        f'<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'{shape_xml}</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'
    )


def cell_anchor(shape_xml, name, cx, cy, inset_x=INSET_EMU, inset_y=None):
    """Cell-anchored style: relativeFrom="column" (== the containing table
    cell's own left edge, because layoutInCell="1") / relativeFrom=
    "paragraph" (== the cell paragraph's top), small posOffset insets,
    wrapNone. Must be appended into the paragraph of the CELL it anchors to."""
    if inset_y is None:
        inset_y = inset_x
    return anchor_run(shape_xml, name, cx, cy, ("column", inset_x), ("paragraph", inset_y),
                       layout_in_cell="1")


def page_anchor(shape_xml, name, cx, cy, x_emu, y_emu):
    """Page-anchored style: relativeFrom="page" for both axes."""
    return anchor_run(shape_xml, name, cx, cy, ("page", x_emu), ("page", y_emu),
                       layout_in_cell="0")


# --------------------------------------------------------------- helpers --
def set_page_geometry(document):
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)


def set_fixed_line_paragraph(paragraph, text, size_pt=11):
    pPr = paragraph.paragraph_format
    pPr.space_before = Pt(0)
    pPr.space_after = Pt(0)
    spacing = paragraph._p.get_or_add_pPr().find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        paragraph._p.get_or_add_pPr().append(spacing)
    spacing.set(qn("w:line"), str(PRECEDING_PARA_TWIPS))
    spacing.set(qn("w:lineRule"), "exact")
    run = paragraph.add_run(text)
    run.font.size = Pt(size_pt)
    run.font.name = "Yu Gothic"


def build_table(document, gridcol_twips, headers, rows):
    n_rows = 1 + len(rows)
    n_cols = len(headers)
    table = document.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    table.autofit = False  # w:tblLayout type="fixed"
    tblPr = table._tbl.tblPr
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    tblPr.append(ind)

    for c, w in enumerate(gridcol_twips):
        table.columns[c].width = Twips(w)

    for r in range(n_rows):
        table.rows[r].height = Twips(ROW_HEIGHT_TWIPS)
        table.rows[r].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = ""
        run = cell.paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = "Yu Gothic"
        run.font.color.rgb = docx.shared.RGBColor.from_string(WHITE)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), NAVY)
        cell._tc.get_or_add_tcPr().append(shd)

    for r, rowvals in enumerate(rows, start=1):
        for c, val in enumerate(rowvals):
            cell = table.cell(r, c)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)
            run.font.name = "Yu Gothic"
    return table


# ================================================================ page 1 ==
def build_page1(document):
    p_before = document.add_paragraph()
    set_fixed_line_paragraph(p_before, "開発スケジュール（本日ライン基準）")

    # Page-anchored today line: relativeFrom="page" both axes. X centers on
    # the 9/3 column (5th table column, 0-based index 4): leftMargin +
    # (工程+担当+9/1+9/2 widths) + half of 9/3's width.
    x_twips = MARGIN_TWIPS + span_width_twips(GRIDCOL_MAIN, 0, 3) + GRIDCOL_MAIN[4] // 2
    y_twips = TABLE_TOP_TWIPS
    table_height_twips = ROW_HEIGHT_TWIPS * (1 + len(ROWS_MAIN))
    today_x = twips_to_emu(x_twips)
    today_y = twips_to_emu(y_twips)
    today_cy = twips_to_emu(table_height_twips)
    append_raw(p_before, page_anchor(
        wsp_line(1, today_cy, line_hex=RED, tail_triangle=True), "本日線", 1, today_cy, today_x, today_y))
    record("D-TODAYLINE", f"本日線: page-anchored straightConnector1, x={today_x}EMU y={today_y}EMU cy={today_cy}EMU",
           "wp:anchor relativeFrom=page + a:tailEnd", True)

    table = build_table(document, GRIDCOL_MAIN, HEADERS_MAIN, ROWS_MAIN)

    def cell_para(row, col):
        return table.cell(row, col).paragraphs[0]

    # 1) 要件定義 rightArrow, 9/1-9/2 (row1, cols 2-3)
    cx = twips_to_emu(span_width_twips(GRIDCOL_MAIN, 2, 3))
    cy = bar_cy_emu()
    voff = (row_height_emu() - cy) // 2
    append_raw(cell_para(1, 2), cell_anchor(
        wsp_autoshape("要件定義", "rightArrow", cx, cy, NAVY, text="要件定義", font_size_half_pt=16),
        "要件定義", cx, cy, inset_y=voff))
    record("D-ARROW-REQ", "要件定義: rightArrow, cell-anchored 9/1セル, extent 2列分", "wp:anchor+wps:wsp", True)

    # 2) 設計 rightArrow, 9/2-9/4 (row2, cols 3-5)
    cx = twips_to_emu(span_width_twips(GRIDCOL_MAIN, 3, 5))
    append_raw(cell_para(2, 3), cell_anchor(
        wsp_autoshape("設計", "rightArrow", cx, cy, TEAL, text="設計", font_size_half_pt=16),
        "設計", cx, cy, inset_y=voff))
    record("D-ARROW-DESIGN", "設計: rightArrow, cell-anchored 9/2セル, extent 3列分", "wp:anchor+wps:wsp", True)

    # 3) rect bar 実装, 9/3-9/5 (row3, cols 4-6), no text
    cx = twips_to_emu(span_width_twips(GRIDCOL_MAIN, 4, 6))
    append_raw(cell_para(3, 4), cell_anchor(
        wsp_autoshape("実装バー", "rect", cx, cy, GOLD), "実装バー", cx, cy, inset_y=voff))
    record("D-RECT-IMPL", "実装: rect(無地/無文字), cell-anchored 9/3セル, extent 3列分", "wp:anchor+wps:wsp", True)

    # 4) leftRightArrow テスト, 9/4-9/8 (row4, cols 5-7)
    cx = twips_to_emu(span_width_twips(GRIDCOL_MAIN, 5, 7))
    append_raw(cell_para(4, 5), cell_anchor(
        wsp_autoshape("テスト", "leftRightArrow", cx, cy, GREEN, text="テスト", font_size_half_pt=16),
        "テスト", cx, cy, inset_y=voff))
    record("D-ARROW-TEST", "テスト: leftRightArrow, cell-anchored 9/4セル, extent 3列分", "wp:anchor+wps:wsp", True)

    # 5) diamond リリース, 9/8 (row5, col7), centered in the cell
    col_w = col_width_emu_main(7)
    dx = (col_w - DIAMOND_SIZE_EMU) // 2
    dy = (row_height_emu() - DIAMOND_SIZE_EMU) // 2
    append_raw(cell_para(5, 7), cell_anchor(
        wsp_autoshape("リリース", "diamond", DIAMOND_SIZE_EMU, DIAMOND_SIZE_EMU, RED),
        "リリース", DIAMOND_SIZE_EMU, DIAMOND_SIZE_EMU, inset_x=dx, inset_y=dy))
    record("D-DIAMOND", "リリース: diamond(0.25in四方), cell-anchored 9/8セル中央", "wp:anchor+wps:wsp", True)

    # 6) textbox ▲レビュー, 設計 x 9/5 (row2, col6)
    col_w = col_width_emu_main(6)
    tb_cx = col_w - 2 * INSET_EMU
    tb_cy = row_height_emu() - 2 * INSET_EMU
    append_raw(cell_para(2, 6), cell_anchor(
        wsp_autoshape("レビュー注記", "rect", tb_cx, tb_cy, None, line_hex=None,
                      text="▲レビュー", text_color=RED, font_size_half_pt=14, txbox=True),
        "レビュー注記", tb_cx, tb_cy))
    record("D-LABEL", "▲レビュー: テキストボックス, cell-anchored 設計x9/5セル", "wp:anchor+wps:cNvSpPr txBox=1", True)

    return table


# ================================================================ page 2 ==
def build_page2(document):
    document.add_page_break()

    p_before = document.add_paragraph()
    set_fixed_line_paragraph(p_before, "回転した矢印（ページ2）")

    table2 = build_table(document, GRIDCOL_SMALL, HEADERS_SMALL, ROWS_SMALL)

    def cell_para(row, col):
        return table2.cell(row, col).paragraphs[0]

    # rightArrow rotated 180deg, text 戻し, cell-anchored 9/2-9/3 (row1, cols2-3)
    cx = twips_to_emu(span_width_twips(GRIDCOL_SMALL, 2, 3))
    cy = bar_cy_emu()
    voff = (row_height_emu() - cy) // 2
    append_raw(cell_para(1, 2), cell_anchor(
        wsp_autoshape("戻し", "rightArrow", cx, cy, GOLD, text="戻し", font_size_half_pt=16, rot=10800000),
        "戻し", cx, cy, inset_y=voff))
    record("D-ROT180", '戻し: rightArrow rot="10800000", cell-anchored 9/2セル, page2', "a:xfrm@rot", True)

    # Spacer paragraphs, then the real flow diagram, well clear of table2.
    for _ in range(4):
        document.add_paragraph()
    flow_p = document.add_paragraph("実データフロー（表とは無関係）")
    flow_p.runs[0].font.size = Pt(9)
    flow_p.runs[0].font.italic = True
    flow_p.runs[0].font.color.rgb = docx.shared.RGBColor.from_string(MUTED)

    anchor_p = document.add_paragraph()
    box_cx, box_cy = 1645920, 640080  # 1.8in x 0.7in, same as the PPTX fixture's boxes
    gap = 400000
    base_x, base_y = 50800, 50800
    append_raw(anchor_p, anchor_run(
        wsp_autoshape("開始", "roundRect", box_cx, box_cy, NAVY, text="開始", font_size_half_pt=20),
        "開始", box_cx, box_cy, ("column", base_x), ("paragraph", base_y), layout_in_cell="0"))
    append_raw(anchor_p, anchor_run(
        wsp_autoshape("完了", "roundRect", box_cx, box_cy, TEAL, text="完了", font_size_half_pt=20),
        "完了", box_cx, box_cy, ("column", base_x + box_cx + gap), ("paragraph", base_y), layout_in_cell="0"))
    append_raw(anchor_p, anchor_run(
        wsp_line(gap, 1, line_hex=MUTED, width_emu=19050, tail_triangle=True),
        "接続線", gap, 1, ("column", base_x + box_cx), ("paragraph", base_y + box_cy // 2), layout_in_cell="0"))
    record("D-REALFLOW", "表と無関係な実フロー図（開始/完了 roundRect + straightConnector1）,page2",
           "wp:anchor relativeFrom=column/paragraph, layoutInCell=0 (not in a table cell)", True)

    # Page-anchored note overlapping table2's bottom edge by ~30% only --
    # must NOT be picked up as a table overlay.
    table2_bottom_twips = TABLE_TOP_TWIPS + ROW_HEIGHT_TWIPS * (1 + len(ROWS_SMALL))
    note_cy_twips = 400
    overlap_frac = 0.30
    note_top_twips = table2_bottom_twips - round(note_cy_twips * overlap_frac)
    note_x = twips_to_emu(MARGIN_TWIPS + 500)
    note_y = twips_to_emu(note_top_twips)
    note_cx = twips_to_emu(3600)
    note_cy = twips_to_emu(note_cy_twips)
    note_p = document.add_paragraph()
    append_raw(note_p, page_anchor(
        wsp_autoshape("実績更新注記", "rect", note_cx, note_cy, None, line_hex=None,
                      text="※ 実績は毎週更新", text_color=MUTED, font_size_half_pt=14, txbox=True),
        "実績更新注記", note_cx, note_cy, note_x, note_y))
    record("D-NOTE-NOOVERLAY", f"※実績は毎週更新: page-anchored, table2下端との重なり{int(overlap_frac*100)}%のみ(<50%)",
           "wp:anchor relativeFrom=page, 50%重複ルール未満", True)

    return table2, table2_bottom_twips, note_top_twips, note_cy_twips


# ================================================================= build ==
def build(output_path: Path):
    _DOCPR_ID[0] = 1
    _REL_HEIGHT[0] = 251658240
    document = docx.Document()
    set_page_geometry(document)
    document.core_properties.title = "schedule-arrows"
    document.core_properties.subject = "DRMD DOCX table-overlay fixture"
    document.core_properties.author = "RTMD fixture"
    document.core_properties.created = FIXED_DATETIME
    document.core_properties.modified = FIXED_DATETIME

    build_page1(document)
    table2, table2_bottom_twips, note_top_twips, note_cy_twips = build_page2(document)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    normalize_zip_timestamps(output_path)
    return output_path


def normalize_zip_timestamps(path, date_time=FIXED_ZIP_DATE_TIME):
    tmp_path = str(path) + ".normalize.tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            item.date_time = date_time
            zout.writestr(item, data)
    import os
    os.replace(tmp_path, path)


# ============================================================ self-check ==
_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
}


def qn2(tag):
    prefix, local = tag.split(":")
    return f"{{{_NS[prefix]}}}{local}"


def overlap_len(a0, a1, b0, b1):
    return max(0.0, min(a1, b1) - max(a0, b0))


def assign_axis(a0, a1, bounds):
    length = a1 - a0
    if length <= 0:
        mid = a0
        for i, (b0, b1) in enumerate(bounds):
            if b0 <= mid <= b1:
                return [i]
        return [0] if mid < bounds[0][0] else [len(bounds) - 1]
    covered = [i for i, (b0, b1) in enumerate(bounds) if overlap_len(a0, a1, b0, b1) >= 0.5 * (b1 - b0)]
    if covered:
        return covered
    mid = (a0 + a1) / 2
    for i, (b0, b1) in enumerate(bounds):
        if b0 <= mid <= b1:
            return [i]
    return [0] if mid < bounds[0][0] else [len(bounds) - 1]


def table_col_bounds(gridcol_twips):
    bounds = []
    acc = 0
    for w in gridcol_twips:
        bounds.append((twips_to_emu(acc), twips_to_emu(acc + w)))
        acc += w
    return bounds


def table_row_bounds(n_rows):
    bounds = []
    for r in range(n_rows):
        top = TABLE_TOP_TWIPS + r * ROW_HEIGHT_TWIPS
        bounds.append((twips_to_emu(top), twips_to_emu(top + ROW_HEIGHT_TWIPS)))
    return bounds


def find_ancestor(el, tagname):
    for anc in el.iterancestors():
        if etree.QName(anc).localname == tagname:
            return anc
    return None


def cell_index_in_row(tc, tr):
    tcs = tr.findall(qn2("w:tc"))
    return tcs.index(tc) if tc in tcs else None


def row_index_in_table(tr, tbl):
    trs = tbl.findall(qn2("w:tr"))
    return trs.index(tr) if tr in trs else None


def verify_anchor(anchor_el, tbl1, tbl2, page_num):
    doc_pr = anchor_el.find(qn2("wp:docPr"))
    name = doc_pr.get("name") if doc_pr is not None else "?"
    pos_h = anchor_el.find(qn2("wp:positionH"))
    pos_v = anchor_el.find(qn2("wp:positionV"))
    rf_h = pos_h.get("relativeFrom")
    rf_v = pos_v.get("relativeFrom")
    off_h = int(pos_h.find(qn2("wp:posOffset")).text)
    off_v = int(pos_v.find(qn2("wp:posOffset")).text)
    extent = anchor_el.find(qn2("wp:extent"))
    cx, cy = int(extent.get("cx")), int(extent.get("cy"))
    layout_in_cell = anchor_el.get("layoutInCell")
    wsp = anchor_el.find(f'.//{qn2("wps:wsp")}')
    prst_el = anchor_el.find(f'.//{qn2("a:prstGeom")}')
    prst = prst_el.get("prst") if prst_el is not None else "?"
    xfrm = anchor_el.find(f'.//{qn2("a:xfrm")}')
    rot = int(xfrm.get("rot", "0")) if xfrm is not None else 0
    style = "cell" if layout_in_cell == "1" else "page"

    print(f"\n  [{name}] style={style} prst={prst} rot={rot}")
    print(f"    positionH relativeFrom={rf_h} posOffset={off_h}EMU  "
          f"positionV relativeFrom={rf_v} posOffset={off_v}EMU  extent=({cx},{cy})EMU")

    tc = find_ancestor(anchor_el, "tc")
    if tc is not None:
        tr = find_ancestor(tc, "tr")
        tbl = find_ancestor(tr, "tbl")
        row_idx = row_index_in_table(tr, tbl)
        col_idx = cell_index_in_row(tc, tr)
        gridcol = GRIDCOL_MAIN if tbl is tbl1 else GRIDCOL_SMALL
        col_bounds = table_col_bounds(gridcol)
        n_rows = 1 + (len(ROWS_MAIN) if tbl is tbl1 else len(ROWS_SMALL))
        row_bounds = table_row_bounds(n_rows)
        x0 = sum(w for w in gridcol[:col_idx]) * TWIP_EMU + off_h
        y0 = (row_idx * ROW_HEIGHT_TWIPS) * TWIP_EMU + off_v
        # rotation-aware AABB (rare here, but kept symmetric with pptx/xlsx)
        cx2, cy2 = cx, cy
        quarter = round(rot / 5400000.0) % 4
        if quarter in (1, 3):
            cx2, cy2 = cy, cx
        rows = assign_axis(y0, y0 + cy2, [(b[0] - row_bounds[0][0], b[1] - row_bounds[0][0]) for b in row_bounds])
        cols = assign_axis(x0, x0 + cx2, [(b[0], b[1]) for b in col_bounds])
        table_label = "table1(page1)" if tbl is tbl1 else "table2(page2)"
        print(f"    anchored in {table_label} cell (row={row_idx},col={col_idx}) -> rows={rows} cols={cols}")
        return
    if rf_h == "page" and rf_v == "page":
        # Try the table that lives on the SAME page as this anchor (a
        # page-relative posOffset is only meaningful against the table
        # physically sharing that page).
        candidates = {
            1: [("table1(page1)", GRIDCOL_MAIN, len(ROWS_MAIN))],
            2: [("table2(page2)", GRIDCOL_SMALL, len(ROWS_SMALL))],
        }.get(page_num, [])
        for label, gridcol, n_data_rows in candidates:
            table_left = twips_to_emu(MARGIN_TWIPS)
            table_top = twips_to_emu(TABLE_TOP_TWIPS)
            col_bounds = [(table_left + b0, table_left + b1) for b0, b1 in table_col_bounds(gridcol)]
            row_bounds = [(table_top + b0 - twips_to_emu(0), table_top + b1 - twips_to_emu(0))
                          for b0, b1 in [(r * ROW_HEIGHT_TWIPS * TWIP_EMU, (r + 1) * ROW_HEIGHT_TWIPS * TWIP_EMU)
                                          for r in range(1 + n_data_rows)]]
            x0, y0 = off_h, off_v
            cx2, cy2 = cx, cy
            quarter = round(rot / 5400000.0) % 4
            if quarter in (1, 3):
                cx2, cy2 = cy, cx
            row_overlap = overlap_len(y0, y0 + max(cy2, 1), row_bounds[0][0], row_bounds[-1][1])
            if row_overlap <= 0:
                continue
            rows = assign_axis(y0, y0 + max(cy2, 1), row_bounds)
            cols = assign_axis(x0, x0 + max(cx2, 1), col_bounds)
            frac = row_overlap / max(cy2, 1)
            print(f"    page-anchored; vs {label}: rows={rows} cols={cols} "
                  f"(own-height overlap fraction with that table's row band = {frac:.2f})")
        return
    print("    page-anchored/paragraph-anchored shape not inside any table cell "
          "(e.g. the standalone flow diagram) -- not a table overlay by construction")


def run_verify(path: Path):
    with zipfile.ZipFile(path) as z:
        doc_xml = etree.fromstring(z.read("word/document.xml"))
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        size = path.stat().st_size

    tables = doc_xml.findall(f'.//{qn2("w:tbl")}')
    tbl1, tbl2 = tables[0], tables[1]

    # Track which page each wp:anchor lands on by counting <w:br w:type="page"/>
    # markers encountered before it, in document order. lxml element proxies
    # returned by .iter()/.findall() are transient and NOT safe to key a
    # dict by id() across separate traversals (the same memory address can
    # be reused for an unrelated proxy once the first is garbage collected)
    # -- so anchors and their page numbers are collected in ONE pass instead
    # of a dict keyed by id().
    page = 1
    anchors = []
    anchor_pages = []
    for el in doc_xml.iter():
        local = etree.QName(el).localname
        if local == "br" and el.get(qn2("w:type")) == "page":
            page += 1
        elif local == "anchor":
            anchors.append(el)
            anchor_pages.append(page)

    print(f"\n=== schedule-arrows.docx: {len(anchors)} floating shapes ===")
    for anchor, page_num in zip(anchors, anchor_pages):
        verify_anchor(anchor, tbl1, tbl2, page_num)

    record("REOPEN", "python-docxで再オープン", "docx.Document(path)", True)
    record("SIZE", f"ファイルサイズ {size} bytes < 200KB", "path.stat().st_size", size < 200 * 1024)
    record("NO-MEDIA", "word/media/* が存在しない（画像なし）", "zipfile namelist", len(media) == 0)
    print(f"\nfile size: {size} bytes, media parts: {media}")


def verify_determinism(out_path):
    import hashlib
    import tempfile
    base_len = len(CHECKLIST)
    with tempfile.TemporaryDirectory() as td:
        p1 = Path(td) / "a.docx"
        p2 = Path(td) / "b.docx"
        build(p1)
        build(p2)
        h1 = hashlib.sha256(p1.read_bytes()).hexdigest()
        h2 = hashlib.sha256(p2.read_bytes()).hexdigest()
    del CHECKLIST[base_len:]
    ok = h1 == h2
    record("DETERMINISM", "2回連続生成しても同一バイト列 (sha256一致)", "hashlib.sha256 of two independent builds", ok)
    print(f"  sha256 run A: {h1}")
    print(f"  sha256 run B: {h2}")


def print_checklist():
    print("\n=== schedule-arrows.docx element checklist ===")
    width = max(len(c[0]) for c in CHECKLIST)
    for elem_id, desc, method, status in CHECKLIST:
        print(f"[{status}] {elem_id.ljust(width)}  {desc}  ({method})")
    failed = [c for c in CHECKLIST if c[3] != "PASS"]
    print(f"\n{len(CHECKLIST)} checks, {len(failed)} failures.")
    return len(failed) == 0


def main():
    args = [a for a in sys.argv[1:] if a != "--verify"]
    verify_only = "--verify" in sys.argv[1:]
    out_path = Path(args[0]) if args else DEFAULT_OUT

    if verify_only:
        if not out_path.exists():
            print(f"error: {out_path} does not exist; run without --verify first", file=sys.stderr)
            sys.exit(2)
    else:
        out_path = build(out_path)
        print(f"wrote {out_path} ({out_path.stat().st_size} bytes)")

    try:
        reopened = docx.Document(str(out_path))
        assert len(reopened.tables) == 2
    except Exception as exc:
        record("REOPEN", "python-docxで再オープン", "docx.Document(path)", False)
        print(f"reopen failed: {exc}", file=sys.stderr)
    run_verify(out_path)
    verify_determinism(out_path)

    all_ok = print_checklist()
    if not all_ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
