#!/usr/bin/env python3
"""DocRedock corpus regression harness (P2: 実業務文書コーパスによる回帰試験).

合成 fixture では拾いきれない「実業務の Word/Excel/PowerPoint/PDF」を大量に投入し、
ファイル単位・集計単位で content retention rate / warning rate / conversion failure
rate を測定し、過去の report.json (--baseline) と突き合わせて回帰を検出するハーネス。

DocRedock.Cli には export/restore の JSON レポートが無いため (--verbose はテキストの
詳細行を増やすだけ)、標準出力のテキスト行 ("WARNING <code>: <message>" 等) を解析する。
`tools/conversion-qa/run.py` の道具立て (dotnet run フォールバック、タイムアウト付き
subprocess 実行、report.json/report.md の 2 本立て) を踏襲する。

標準ライブラリのみで動作する。python-docx / python-pptx / openpyxl / pypdf が無い
環境では、該当フォーマットの content retention だけを null にして警告し、それ以外の
指標 (warning rate / failure rate / wall time / roundtrip identity) は測定を続ける。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from typing import Optional

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CLI_BINARY = REPO_ROOT / "src" / "DocRedock.Cli" / "bin" / "Release" / "net10.0" / "DocRedock.Cli"
DEFAULT_OUT_ROOT = REPO_ROOT / "artifacts" / "corpus-regression"
SUPPORTED_SUFFIXES = {".docx", ".xlsx", ".pptx", ".pdf"}
SCHEMA_VERSION = "1"

DEFAULT_TIMEOUT_SEC = 300
DEFAULT_TOLERANCE_PCT_POINTS = 1.0
STDOUT_TAIL_CHARS = 2000


class CliResolutionError(RuntimeError):
    pass


class ExtractorUnavailable(Exception):
    """Raised when the optional extraction library for a format is not installed."""


# --------------------------------------------------------------------------
# Tokenizer: multiset of "content tokens" for retention / spurious scoring.
#
#   - CJK (hiragana/katakana/kanji): character bigrams over each contiguous
#     CJK run; a lone character becomes a 1-character token.
#   - Latin/other alphabetic scripts: the whole contiguous run, lower-cased,
#     is one "word" token.
#   - Digits: the whole contiguous digit run is one token.
#   - Everything else (Markdown/DRMD punctuation, whitespace) is a separator
#     and is never itself a token.
# --------------------------------------------------------------------------

_CJK_RANGES = (
    (0x3040, 0x30FF),   # Hiragana + Katakana
    (0x3400, 0x4DBF),   # CJK Unified Ideographs Extension A
    (0x4E00, 0x9FFF),   # CJK Unified Ideographs
    (0xF900, 0xFAFF),   # CJK Compatibility Ideographs
    (0xFF66, 0xFF9F),   # Halfwidth Katakana
)


def _is_cjk(ch: str) -> bool:
    code = ord(ch)
    return any(lo <= code <= hi for lo, hi in _CJK_RANGES)


def _classify_char(ch: str) -> str:
    if ch.isdigit():
        return "digit"
    if _is_cjk(ch):
        return "cjk"
    if ch.isalpha():
        return "latin"
    return "other"


def tokenize(text: str) -> list:
    """Splits text into content tokens (see module docstring for the rules)."""
    tokens: list = []
    if not text:
        return tokens
    for word in text.split():
        length = len(word)
        index = 0
        while index < length:
            cls = _classify_char(word[index])
            end = index + 1
            while end < length and _classify_char(word[end]) == cls:
                end += 1
            run = word[index:end]
            if cls == "cjk":
                if len(run) == 1:
                    tokens.append(run)
                else:
                    tokens.extend(run[k:k + 2] for k in range(len(run) - 1))
            elif cls == "latin":
                tokens.append(run.lower())
            elif cls == "digit":
                tokens.append(run)
            # "other" (punctuation/symbols) is a separator, never a token.
            index = end
    return tokens


# --------------------------------------------------------------------------
# Markdown / DRMD control-comment stripping (best-effort, not a full parser).
# --------------------------------------------------------------------------

_FRONT_MATTER_RE = re.compile(r"\A---\r?\n.*?\r?\n---\r?\n", re.DOTALL)
_HTML_COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)
_CODE_FENCE_RE = re.compile(r"^```[^\n]*$", re.MULTILINE)
_LINK_URL_RE = re.compile(r"\]\([^)]*\)")


def strip_markdown(text: str) -> str:
    """Removes DRMD front matter/control comments and Markdown link URLs.

    Most remaining Markdown punctuation (`#`, `*`, `_`, `` ` ``, `|`, `-`, `>`)
    is already dropped by :func:`tokenize`'s "other" bucket, so this function
    only needs to remove the pieces that would otherwise leak non-content text
    (structural metadata, file paths) into the token multiset.
    """
    text = _FRONT_MATTER_RE.sub("", text, count=1)
    text = _HTML_COMMENT_RE.sub(" ", text)
    text = _CODE_FENCE_RE.sub("", text)
    text = _LINK_URL_RE.sub("]", text)
    return text


# --------------------------------------------------------------------------
# Content retention / spurious-token scoring.
# --------------------------------------------------------------------------

def compute_content_metrics(source_text: str, output_markdown: str) -> dict:
    source_tokens = Counter(tokenize(source_text))
    output_tokens = Counter(tokenize(strip_markdown(output_markdown)))
    source_total = sum(source_tokens.values())
    output_total = sum(output_tokens.values())

    intersection = sum(min(count, output_tokens.get(token, 0)) for token, count in source_tokens.items())
    retention = (intersection / source_total) if source_total else None

    spurious_count = sum(max(0, count - source_tokens.get(token, 0)) for token, count in output_tokens.items())
    spurious_rate = (spurious_count / output_total) if output_total else None

    missing = [
        (token, count - output_tokens.get(token, 0))
        for token, count in source_tokens.items()
        if count - output_tokens.get(token, 0) > 0
    ]
    missing.sort(key=lambda item: (-item[1], item[0]))

    return {
        "retention": retention,
        "spurious_rate": spurious_rate,
        "source_token_count": source_total,
        "output_token_count": output_total,
        "missing_top_tokens": [{"token": token, "missing_count": count} for token, count in missing[:20]],
    }


# --------------------------------------------------------------------------
# Source text extraction (optional third-party libraries; degrade gracefully).
# --------------------------------------------------------------------------

def extract_docx_text(path: Path) -> str:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise ExtractorUnavailable("python-docx is not installed") from exc

    document = Document(str(path))
    parts: list = []

    def collect_table(table) -> None:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
                for nested in cell.tables:
                    collect_table(nested)

    parts.extend(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        collect_table(table)
    for section in document.sections:
        for holder in (section.header, section.footer):
            try:
                parts.extend(paragraph.text for paragraph in holder.paragraphs)
            except Exception:  # noqa: BLE001 - python-docx raises various things for unlinked headers
                pass
    return "\n".join(parts)


def extract_xlsx_text(path: Path) -> str:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise ExtractorUnavailable("openpyxl is not installed") from exc

    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts: list = []
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        parts.append(str(cell.value))
    finally:
        workbook.close()
    return "\n".join(parts)


def extract_pptx_text(path: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise ExtractorUnavailable("python-pptx is not installed") from exc

    presentation = Presentation(str(path))
    parts: list = []

    def walk(shapes) -> None:
        for shape in shapes:
            sub_shapes = getattr(shape, "shapes", None)
            if sub_shapes is not None:
                try:
                    walk(sub_shapes)
                    continue
                except TypeError:
                    pass
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)

    for slide in presentation.slides:
        walk(slide.shapes)
        if getattr(slide, "has_notes_slide", False) and slide.notes_slide is not None:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                parts.append(notes_frame.text)
    return "\n".join(parts)


def extract_pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as exc:
        raise ExtractorUnavailable("pypdf is not installed") from exc

    reader = PdfReader(str(path))
    parts: list = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - pypdf raises assorted errors on malformed pages
            continue
    return "\n".join(parts)


EXTRACTORS = {
    ".docx": extract_docx_text,
    ".xlsx": extract_xlsx_text,
    ".pptx": extract_pptx_text,
    ".pdf": extract_pdf_text,
}


# --------------------------------------------------------------------------
# CLI discovery and invocation.
# --------------------------------------------------------------------------

def _is_executable(path: Path) -> bool:
    return path.is_file() and os.access(path, os.X_OK)


def resolve_cli(cli_arg: Optional[str]) -> tuple:
    """Returns (command_prefix, kind); kind is "binary" or "dotnet-run"."""
    if cli_arg:
        candidate = Path(cli_arg).expanduser()
        if _is_executable(candidate):
            return [str(candidate)], "binary"
        if cli_arg.strip().lower() in ("dotnet", "dotnet-run"):
            dotnet = shutil.which("dotnet") or str(REPO_ROOT / ".tmp" / "dotnet" / "dotnet")
            if not (shutil.which("dotnet") or _is_executable(Path(dotnet))):
                raise CliResolutionError("--cli dotnet was requested but 'dotnet' was not found on PATH.")
            return [dotnet, "run", "--project", "src/DocRedock.Cli", "-c", "Release", "--"], "dotnet-run"
        raise CliResolutionError(f"--cli {cli_arg!r} is not an executable file, and is not 'dotnet'.")

    if _is_executable(DEFAULT_CLI_BINARY):
        return [str(DEFAULT_CLI_BINARY)], "binary"

    dotnet = shutil.which("dotnet")
    if dotnet is None:
        fallback = REPO_ROOT / ".tmp" / "dotnet" / "dotnet"
        if _is_executable(fallback):
            dotnet = str(fallback)
    if dotnet is None:
        raise CliResolutionError(
            f"Release binary not found at {DEFAULT_CLI_BINARY}, and 'dotnet' is not on PATH; "
            "build the CLI or pass --cli explicitly."
        )
    return [dotnet, "run", "--project", "src/DocRedock.Cli", "-c", "Release", "--"], "dotnet-run"


def build_once(prefix: list, kind: str) -> tuple:
    """Turns a `dotnet run` prefix into one build plus direct execution of the built DLL.

    `dotnet run` evaluates and up-to-date-checks the whole project graph on every call and leaves
    MSBuild worker nodes resident; per corpus file - and in parallel under --jobs - that dominated
    this harness's memory and time. The build disables node reuse and the compiler server so it
    leaves nothing running. If it fails, the original `dotnet run` prefix is kept."""
    if kind != "dotnet-run":
        return prefix, kind
    dotnet = prefix[0]
    try:
        build = subprocess.run(
            [dotnet, "build", "src/DocRedock.Cli", "-c", "Release", "--nologo", "-v", "q",
             "-m:4", "-nodeReuse:false", "-p:UseSharedCompilation=false"],
            cwd=str(REPO_ROOT), capture_output=True, text=True, timeout=900,
        )
    except (OSError, subprocess.TimeoutExpired):
        return prefix, kind
    built = sorted((REPO_ROOT / "src" / "DocRedock.Cli" / "bin" / "Release").glob("*/DocRedock.Cli.dll"),
                   key=lambda path: path.stat().st_mtime, reverse=True)
    if build.returncode != 0 or not built:
        return prefix, kind
    return [dotnet, str(built[0])], "dotnet-built-once"


def run_cli(cmd: list, timeout: int) -> dict:
    start = time.time()
    environment = os.environ.copy()
    # PDF export and the roundtrip/restore profile are experimental; the harness
    # is meaningless without them, so force the gate on regardless of the
    # caller's ambient environment.
    environment["DOCREDOCK_ENABLE_EXPERIMENTAL"] = "1"
    try:
        proc = subprocess.run(
            cmd, cwd=str(REPO_ROOT), capture_output=True, text=True,
            timeout=timeout, env=environment,
        )
        return {
            "command": cmd, "returncode": proc.returncode, "stdout": proc.stdout,
            "stderr": proc.stderr, "elapsed_sec": round(time.time() - start, 3), "timed_out": False,
        }
    except subprocess.TimeoutExpired as exc:
        # TimeoutExpired retains bytes even when subprocess.run uses text=True;
        # a timeout can also split a UTF-8 character in the captured output.
        stdout = exc.stdout.decode("utf-8", errors="replace") if isinstance(exc.stdout, bytes) else exc.stdout or ""
        stderr = exc.stderr.decode("utf-8", errors="replace") if isinstance(exc.stderr, bytes) else exc.stderr or ""
        return {
            "command": cmd, "returncode": None, "stdout": stdout,
            "stderr": stderr + f"\n[corpus-regression] timed out after {timeout}s",
            "elapsed_sec": round(time.time() - start, 3), "timed_out": True,
        }
    except OSError as exc:
        return {
            "command": cmd, "returncode": None, "stdout": "",
            "stderr": f"[corpus-regression] failed to launch: {exc}",
            "elapsed_sec": round(time.time() - start, 3), "timed_out": False,
        }


def classify_status(returncode: Optional[int], timed_out: bool) -> str:
    if timed_out:
        return "timeout"
    if returncode is None:
        return "error"
    if returncode == 0:
        return "success"
    if returncode == 1:
        return "success_with_warnings"
    return "failed"


def is_failure_status(status: str) -> bool:
    return status in ("failed", "timeout", "error")


# CLI diagnostic lines look like "WARNING <Code>: <message>" (see
# WriteDiagnosticsAsync / WriteMarkdownDiagnostics in CliApplication.cs); an
# optional "(N diagnostic entries; use --verbose for details)" suffix is kept
# as part of the message.
_DIAG_LINE_RE = re.compile(r"^(WARNING|ERROR|INFORMATION)\s+(\S+):\s?(.*)$")
_VISUAL_SUMMARY_RE = re.compile(
    r"^Visual summary: diagrams=(?P<diagrams>\d+); vector_pages=(?P<vector_pages>\d+); "
    r"native=(?P<native>\d+); high=(?P<high>\d+); medium=(?P<medium>\d+); "
    r"unresolved=(?P<unresolved>\d+); fallback=(?P<fallback>\d+); rejected=(?P<rejected>\d+)$",
    re.MULTILINE,
)


def parse_diagnostics(text: str) -> list:
    diagnostics = []
    for line in text.splitlines():
        match = _DIAG_LINE_RE.match(line.strip())
        if match:
            diagnostics.append({"severity": match.group(1), "code": match.group(2), "message": match.group(3)})
    return diagnostics


def parse_visual_summary(stdout: str) -> Optional[dict]:
    match = _VISUAL_SUMMARY_RE.search(stdout)
    if not match:
        return None
    return {key: int(value) for key, value in match.groupdict().items()}


def export_file(cli_prefix: list, source: Path, profile: str, output: Path, ocr_mode: str, timeout: int) -> dict:
    output.parent.mkdir(parents=True, exist_ok=True)
    args = ["export", str(source), "--profile", profile, "--output", str(output),
            "--force", "--quiet", "--ocr", ocr_mode]
    if profile == "roundtrip":
        args += ["--sidecar", "dir"]
    result = run_cli(cli_prefix + args, timeout=timeout)
    result["status"] = classify_status(result["returncode"], result["timed_out"])
    result["diagnostics"] = parse_diagnostics(result["stdout"])
    result["warning_codes"] = sorted({item["code"] for item in result["diagnostics"] if item["severity"] == "WARNING"})
    visual_summary = parse_visual_summary(result["stdout"])
    if visual_summary is not None:
        result["visual_summary"] = visual_summary
    result["output_exists"] = output.is_file()
    return result


def restore_file(cli_prefix: list, markdown: Path, destination: Path, timeout: int) -> dict:
    destination.parent.mkdir(parents=True, exist_ok=True)
    args = ["restore", str(markdown), "--output", str(destination), "--force"]
    result = run_cli(cli_prefix + args, timeout=timeout)
    result["status"] = classify_status(result["returncode"], result["timed_out"])
    # restore writes WARNING/INFORMATION to stdout and ERROR to stderr (see
    # RestoreAsync in CliApplication.cs); scan both streams for diagnostic lines.
    result["diagnostics"] = parse_diagnostics(result["stdout"] + "\n" + result["stderr"])
    result["output_exists"] = destination.is_file()
    return result


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# --------------------------------------------------------------------------
# Corpus discovery.
# --------------------------------------------------------------------------

def discover_corpus_files(root: Path) -> list:
    results = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(name for name in dirnames if not name.startswith("."))
        for name in filenames:
            if name.startswith("~$"):
                continue
            candidate = Path(dirpath) / name
            if candidate.suffix.lower() in SUPPORTED_SUFFIXES:
                results.append(candidate)
    results.sort(key=lambda item: str(item.relative_to(root)).lower())
    return results


# --------------------------------------------------------------------------
# Per-file processing.
# --------------------------------------------------------------------------

def _slim_profile_result(result: dict, redact: bool) -> dict:
    slim = {
        "status": result["status"],
        "returncode": result["returncode"],
        "elapsed_sec": result["elapsed_sec"],
        "timed_out": result["timed_out"],
        "diagnostics": result.get("diagnostics", []),
        "warning_codes": result.get("warning_codes", []),
        "output_exists": result.get("output_exists"),
    }
    for optional_key in ("visual_summary", "content", "restore"):
        if optional_key in result:
            slim[optional_key] = result[optional_key]
    if not redact:
        # stdout/stderr can carry the absolute source path (e.g. in a fatal
        # exception message); only keep them when the report is not meant to
        # be shared verbatim.
        slim["stdout_tail"] = result.get("stdout", "")[-STDOUT_TAIL_CHARS:]
        slim["stderr_tail"] = result.get("stderr", "")[-STDOUT_TAIL_CHARS:]
    return slim


def process_file(cli_prefix: list, corpus_root: Path, path: Path, out_dir: Path,
                  profiles: list, ocr_mode: str, timeout: int, redact: bool, mapping: dict) -> dict:
    relpath = path.relative_to(corpus_root).as_posix()
    file_id = hashlib.sha256(relpath.encode("utf-8")).hexdigest()[:16]
    mapping[file_id] = relpath
    display_name = file_id if redact else relpath

    record: dict = {
        "file_id": file_id,
        "display_name": display_name,
        "format": path.suffix.lower().lstrip("."),
        "size_bytes": path.stat().st_size,
        "profiles": {},
    }

    if "readable" in profiles:
        readable_md = out_dir / "export" / "readable" / f"{relpath}.md"
        result = export_file(cli_prefix, path, "readable", readable_md, ocr_mode, timeout)
        content: Optional[dict] = None
        if result["status"] in ("success", "success_with_warnings") and result["output_exists"]:
            try:
                output_text = readable_md.read_text(encoding="utf-8", errors="replace")
            except OSError as exc:
                content = {"retention": None, "spurious_rate": None, "note": f"failed to read output: {exc}"}
                output_text = None
            if output_text is not None:
                try:
                    source_text = EXTRACTORS[path.suffix.lower()](path)
                    content = compute_content_metrics(source_text, output_text)
                except ExtractorUnavailable as exc:
                    content = {"retention": None, "spurious_rate": None, "note": str(exc)}
                except Exception as exc:  # noqa: BLE001 - real-world files can be malformed/atypical;
                    # a source-extraction failure must degrade this one metric, not crash the harness.
                    content = {"retention": None, "spurious_rate": None,
                               "note": f"source text extraction failed: {type(exc).__name__}: {exc}"}
        else:
            content = {"retention": None, "spurious_rate": None, "note": "export did not succeed; content not evaluated"}
        result["content"] = content
        record["profiles"]["readable"] = _slim_profile_result(result, redact)

    if "roundtrip" in profiles:
        roundtrip_md = out_dir / "export" / "roundtrip" / f"{relpath}.md"
        result = export_file(cli_prefix, path, "roundtrip", roundtrip_md, ocr_mode, timeout)
        restore_info: dict = {"attempted": False, "status": "not-attempted", "f0_identical": None}
        if result["status"] in ("success", "success_with_warnings") and result["output_exists"]:
            restored_path = out_dir / "restore" / f"{relpath}-restored{path.suffix}"
            restore_result = restore_file(cli_prefix, roundtrip_md, restored_path, timeout)
            restore_info = {
                "attempted": True,
                "status": restore_result["status"],
                "returncode": restore_result["returncode"],
                "elapsed_sec": restore_result["elapsed_sec"],
                "warning_codes": sorted({item["code"] for item in restore_result["diagnostics"]
                                         if item["severity"] == "WARNING"}),
                "f0_identical": None,
            }
            if restore_result["status"] in ("success", "success_with_warnings") and restore_result["output_exists"]:
                try:
                    restore_info["f0_identical"] = sha256(path) == sha256(restored_path)
                except OSError as exc:
                    restore_info["note"] = f"hash comparison failed: {exc}"
            else:
                restore_info["note"] = "restore did not produce output"
        else:
            restore_info["note"] = "roundtrip export did not succeed; restore skipped"
        result["restore"] = restore_info
        record["profiles"]["roundtrip"] = _slim_profile_result(result, redact)

    statuses = [profile_result.get("status") for profile_result in record["profiles"].values()]
    statuses.extend(profile_result.get("restore", {}).get("status")
                    for profile_result in record["profiles"].values())
    record["failed"] = any(is_failure_status(status) for status in statuses)
    warning_codes: set = set()
    for profile_result in record["profiles"].values():
        warning_codes.update(profile_result.get("warning_codes", []))
        warning_codes.update(profile_result.get("restore", {}).get("warning_codes", []))
    record["warning_codes"] = sorted(warning_codes)
    record["has_warning"] = bool(warning_codes)
    return record


# --------------------------------------------------------------------------
# Aggregation: overall + per-format failure/warning/retention/timing.
# --------------------------------------------------------------------------

def percentile(values: list, fraction: float) -> Optional[float]:
    if not values:
        return None
    ordered = sorted(values)
    if len(ordered) == 1:
        return ordered[0]
    rank = (len(ordered) - 1) * fraction
    lower = math.floor(rank)
    upper = math.ceil(rank)
    if lower == upper:
        return ordered[int(rank)]
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (rank - lower)


def _aggregate_bucket(files: list) -> dict:
    count = len(files)
    if count == 0:
        return {"count": 0}

    failed = sum(1 for item in files if item["failed"])
    warned = sum(1 for item in files if item["has_warning"])

    def _values(profile: str, *keys):
        collected = []
        for item in files:
            node = item["profiles"].get(profile, {})
            for key in keys:
                node = node.get(key, {}) if isinstance(node, dict) else {}
            if isinstance(node, (int, float)):
                collected.append(node)
        return collected

    retentions = [value for value in _values("readable", "content", "retention") if value is not None]
    spurious = [value for value in _values("readable", "content", "spurious_rate") if value is not None]
    f0_flags = [
        item["profiles"].get("roundtrip", {}).get("restore", {}).get("f0_identical")
        for item in files
    ]
    f0_flags = [flag for flag in f0_flags if flag is not None]
    readable_times = [
        item["profiles"]["readable"]["elapsed_sec"]
        for item in files if "readable" in item["profiles"] and item["profiles"]["readable"].get("elapsed_sec") is not None
    ]
    roundtrip_times = [
        item["profiles"]["roundtrip"]["elapsed_sec"]
        for item in files if "roundtrip" in item["profiles"] and item["profiles"]["roundtrip"].get("elapsed_sec") is not None
    ]

    return {
        "count": count,
        "failure_rate": failed / count,
        "warning_rate": warned / count,
        "retention_mean": (sum(retentions) / len(retentions)) if retentions else None,
        "retention_sample_count": len(retentions),
        "spurious_mean": (sum(spurious) / len(spurious)) if spurious else None,
        "spurious_sample_count": len(spurious),
        "f0_identical_rate": (sum(1 for flag in f0_flags if flag) / len(f0_flags)) if f0_flags else None,
        "f0_identical_sample_count": len(f0_flags),
        "wall_time_readable_sec": {
            "median": percentile(readable_times, 0.5), "p95": percentile(readable_times, 0.95), "count": len(readable_times),
        },
        "wall_time_roundtrip_sec": {
            "median": percentile(roundtrip_times, 0.5), "p95": percentile(roundtrip_times, 0.95), "count": len(roundtrip_times),
        },
    }


def aggregate_results(files: list) -> dict:
    overall = _aggregate_bucket(files)
    by_format = {}
    for fmt in sorted({item["format"] for item in files}):
        by_format[fmt] = _aggregate_bucket([item for item in files if item["format"] == fmt])
    histogram: Counter = Counter()
    for item in files:
        histogram.update(item.get("warning_codes", []))
    return {
        "file_count": len(files),
        "overall": overall,
        "by_format": by_format,
        "warning_code_histogram": dict(sorted(histogram.items(), key=lambda kv: (-kv[1], kv[0]))),
    }


# --------------------------------------------------------------------------
# Baseline comparison / regression detection.
# --------------------------------------------------------------------------

_DELTA_METRICS = {
    # metric -> direction in which an increase is an improvement
    "retention_mean": "higher_better",
    "spurious_mean": "lower_better",
    "warning_rate": "lower_better",
    "failure_rate": "lower_better",
    "f0_identical_rate": "higher_better",
}


def _delta_bucket(current: dict, baseline: dict) -> dict:
    deltas = {}
    for metric in _DELTA_METRICS:
        current_value = current.get(metric)
        baseline_value = baseline.get(metric)
        delta = (current_value - baseline_value) if (current_value is not None and baseline_value is not None) else None
        deltas[metric] = {"current": current_value, "baseline": baseline_value, "delta": delta}
    return deltas


def compute_baseline_comparison(current_report: dict, baseline_report: dict, tolerance: float) -> dict:
    current_aggregate = current_report["aggregate"]
    baseline_aggregate = baseline_report.get("aggregate", {})

    overall_delta = _delta_bucket(current_aggregate.get("overall", {}), baseline_aggregate.get("overall", {}))
    by_format_delta = {}
    all_formats = set(current_aggregate.get("by_format", {})) | set(baseline_aggregate.get("by_format", {}))
    for fmt in sorted(all_formats):
        by_format_delta[fmt] = _delta_bucket(
            current_aggregate.get("by_format", {}).get(fmt, {}),
            baseline_aggregate.get("by_format", {}).get(fmt, {}),
        )

    aggregate_regressions = []

    def _check(scope_label: str, delta_bucket: dict) -> None:
        for metric, direction in _DELTA_METRICS.items():
            entry = delta_bucket.get(metric, {})
            delta = entry.get("delta")
            if delta is None:
                continue
            delta_pct_points = delta * 100
            regressed = (delta_pct_points < -tolerance) if direction == "higher_better" else (delta_pct_points > tolerance)
            if regressed:
                aggregate_regressions.append({
                    "scope": scope_label, "metric": metric,
                    "delta_pct_points": round(delta_pct_points, 3),
                    "current": entry.get("current"), "baseline": entry.get("baseline"),
                })

    _check("overall", overall_delta)
    for fmt, bucket in by_format_delta.items():
        _check(f"format:{fmt}", bucket)

    baseline_by_id = {item["file_id"]: item for item in baseline_report.get("files", [])}
    file_regressions = []
    for item in current_report.get("files", []):
        baseline_item = baseline_by_id.get(item["file_id"])
        if baseline_item is None:
            continue

        current_retention = item["profiles"].get("readable", {}).get("content", {}).get("retention")
        baseline_retention = baseline_item.get("profiles", {}).get("readable", {}).get("content", {}).get("retention")
        if current_retention is not None and baseline_retention is not None:
            drop_pct_points = (baseline_retention - current_retention) * 100
            if drop_pct_points > tolerance:
                file_regressions.append({
                    "file_id": item["file_id"], "display_name": item["display_name"],
                    "kind": "retention_drop", "delta_pct_points": round(drop_pct_points, 3),
                })

        if not baseline_item.get("failed") and item.get("failed"):
            file_regressions.append({
                "file_id": item["file_id"], "display_name": item["display_name"], "kind": "new_failure",
            })

        new_codes = set(item.get("warning_codes", [])) - set(baseline_item.get("warning_codes", []))
        for code in sorted(new_codes):
            file_regressions.append({
                "file_id": item["file_id"], "display_name": item["display_name"],
                "kind": "new_warning_code", "code": code,
            })

    return {
        "aggregate_deltas": {"overall": overall_delta, "by_format": by_format_delta},
        "aggregate_regressions": aggregate_regressions,
        "file_regressions": file_regressions,
        "has_regression": bool(aggregate_regressions or file_regressions),
    }


# --------------------------------------------------------------------------
# Report writers.
# --------------------------------------------------------------------------

def write_report_json(report: dict, out_dir: Path) -> None:
    path = out_dir / "report.json"
    with path.open("w", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2, sort_keys=False)
        handle.write("\n")


def _fmt_pct(value: Optional[float]) -> str:
    return "n/a" if value is None else f"{value * 100:.2f}%"


def _fmt_sec(value: Optional[float]) -> str:
    return "n/a" if value is None else f"{value:.2f}s"


def _bucket_row(label: str, bucket: dict) -> str:
    if bucket.get("count", 0) == 0:
        return f"| {label} | 0 | - | - | - | - | - | - | - |"
    return (
        f"| {label} | {bucket['count']} | {_fmt_pct(bucket.get('failure_rate'))} | "
        f"{_fmt_pct(bucket.get('warning_rate'))} | {_fmt_pct(bucket.get('retention_mean'))} | "
        f"{_fmt_pct(bucket.get('spurious_mean'))} | {_fmt_pct(bucket.get('f0_identical_rate'))} | "
        f"{_fmt_sec(bucket.get('wall_time_readable_sec', {}).get('median'))} | "
        f"{_fmt_sec(bucket.get('wall_time_readable_sec', {}).get('p95'))} |"
    )


def render_report_md(report: dict) -> str:
    lines = ["# corpus-regression report", ""]
    lines.append(f"- schema_version: {report['schema_version']}")
    lines.append(f"- generated_at: {report['generated_at']}")
    lines.append(f"- files: {report['corpus']['file_count']}")
    lines.append(f"- redact: {report['args'].get('redact')}")
    lines.append(f"- cli: `{' '.join(report['tooling']['cli'])}` ({report['tooling']['cli_kind']})")
    lines.append("")

    aggregate = report["aggregate"]
    lines.append("## 集計 (aggregate)")
    lines.append("")
    lines.append("| scope | files | failure rate | warning rate | retention | spurious | F0 identical | median(readable) | p95(readable) |")
    lines.append("| --- | --- | --- | --- | --- | --- | --- | --- | --- |")
    lines.append(_bucket_row("overall", aggregate["overall"]))
    for fmt, bucket in sorted(aggregate["by_format"].items()):
        lines.append(_bucket_row(fmt, bucket))
    lines.append("")

    lines.append("## warning code histogram")
    lines.append("")
    if aggregate["warning_code_histogram"]:
        lines.append("| code | files |")
        lines.append("| --- | --- |")
        for code, count in list(aggregate["warning_code_histogram"].items())[:30]:
            lines.append(f"| {code} | {count} |")
    else:
        lines.append("(warning なし)")
    lines.append("")

    failed_files = [item for item in report["files"] if item["failed"]]
    lines.append(f"## 失敗ファイル ({len(failed_files)})")
    lines.append("")
    if failed_files:
        lines.append("| file | format | readable | roundtrip |")
        lines.append("| --- | --- | --- | --- |")
        for item in failed_files[:100]:
            readable_status = item["profiles"].get("readable", {}).get("status", "-")
            roundtrip_status = item["profiles"].get("roundtrip", {}).get("status", "-")
            lines.append(f"| {item['display_name']} | {item['format']} | {readable_status} | {roundtrip_status} |")
        if len(failed_files) > 100:
            lines.append(f"| ... | ... | ... | (+{len(failed_files) - 100} more; see report.json) |")
    else:
        lines.append("(失敗ファイルはありません)")
    lines.append("")

    comparison = report.get("baseline_comparison")
    if comparison:
        lines.append("## baseline との差分")
        lines.append("")
        lines.append(f"- baseline: `{comparison['baseline_path']}`")
        lines.append(f"- tolerance: {comparison['tolerance_pct_points']} pt")
        lines.append(f"- has_regression: **{comparison['has_regression']}**")
        lines.append("")
        lines.append("| scope | metric | current | baseline | delta(pt) |")
        lines.append("| --- | --- | --- | --- | --- |")

        def _delta_rows(scope: str, bucket: dict) -> None:
            for metric, entry in bucket.items():
                if entry.get("delta") is None:
                    continue
                lines.append(
                    f"| {scope} | {metric} | {_fmt_pct(entry['current'])} | {_fmt_pct(entry['baseline'])} | "
                    f"{entry['delta'] * 100:+.3f} |"
                )

        _delta_rows("overall", comparison["aggregate_deltas"]["overall"])
        for fmt, bucket in sorted(comparison["aggregate_deltas"]["by_format"].items()):
            _delta_rows(f"format:{fmt}", bucket)
        lines.append("")

        if comparison["aggregate_regressions"]:
            lines.append("### aggregate regressions")
            lines.append("")
            for item in comparison["aggregate_regressions"]:
                lines.append(
                    f"- {item['scope']} / {item['metric']}: {item['delta_pct_points']:+.3f} pt "
                    f"(current={_fmt_pct(item['current'])}, baseline={_fmt_pct(item['baseline'])})"
                )
            lines.append("")

        if comparison["file_regressions"]:
            lines.append(f"### per-file regressions ({len(comparison['file_regressions'])})")
            lines.append("")
            for item in comparison["file_regressions"][:100]:
                if item["kind"] == "retention_drop":
                    lines.append(f"- {item['display_name']}: retention drop {item['delta_pct_points']:+.3f} pt")
                elif item["kind"] == "new_failure":
                    lines.append(f"- {item['display_name']}: newly failing")
                else:
                    lines.append(f"- {item['display_name']}: new warning code {item['code']}")
            lines.append("")
    return "\n".join(lines)


def write_report_md(report: dict, out_dir: Path) -> None:
    (out_dir / "report.md").write_text(render_report_md(report), encoding="utf-8")


def print_summary(report: dict) -> None:
    aggregate = report["aggregate"]
    print("[corpus-regression] --- summary ---")
    overall = aggregate["overall"]
    print(
        f"[corpus-regression] overall: files={overall.get('count', 0)} "
        f"failure_rate={_fmt_pct(overall.get('failure_rate'))} "
        f"warning_rate={_fmt_pct(overall.get('warning_rate'))} "
        f"retention={_fmt_pct(overall.get('retention_mean'))} "
        f"spurious={_fmt_pct(overall.get('spurious_mean'))} "
        f"f0_identical={_fmt_pct(overall.get('f0_identical_rate'))}"
    )
    for fmt, bucket in sorted(aggregate["by_format"].items()):
        print(
            f"[corpus-regression]   {fmt}: files={bucket.get('count', 0)} "
            f"failure_rate={_fmt_pct(bucket.get('failure_rate'))} "
            f"warning_rate={_fmt_pct(bucket.get('warning_rate'))} "
            f"retention={_fmt_pct(bucket.get('retention_mean'))} "
            f"spurious={_fmt_pct(bucket.get('spurious_mean'))} "
            f"f0_identical={_fmt_pct(bucket.get('f0_identical_rate'))} "
            f"median={_fmt_sec(bucket.get('wall_time_readable_sec', {}).get('median'))} "
            f"p95={_fmt_sec(bucket.get('wall_time_readable_sec', {}).get('p95'))}"
        )
    comparison = report.get("baseline_comparison")
    if comparison:
        print(f"[corpus-regression] baseline comparison: has_regression={comparison['has_regression']}")
        for item in comparison["aggregate_regressions"]:
            print(f"[corpus-regression]   REGRESSION {item['scope']}/{item['metric']}: {item['delta_pct_points']:+.3f} pt")
        if comparison["file_regressions"]:
            print(f"[corpus-regression]   {len(comparison['file_regressions'])} per-file regression(s); see report.md")


# --------------------------------------------------------------------------
# CLI entry point.
# --------------------------------------------------------------------------

def _timestamp() -> str:
    return datetime.now().strftime("%Y%m%dT%H%M%S")


def parse_args(argv: Optional[list] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="DocRedock corpus regression harness: measure content retention / warning rate / "
                     "failure rate over a real-world document corpus, and diff against a previous run.",
    )
    parser.add_argument("--corpus", required=True, help="ファイルを再帰的に探索するディレクトリ (docx/xlsx/pptx/pdf)")
    parser.add_argument("--out", help="出力先ディレクトリ (既定: artifacts/corpus-regression/<timestamp>/)")
    parser.add_argument("--cli", help="DocRedock.Cli のパス、または 'dotnet' (dotnet run にフォールバック)")
    parser.add_argument("--profiles", default="readable,roundtrip", help="カンマ区切り: readable,roundtrip (既定: 両方)")
    parser.add_argument("--ocr", choices=("on", "off"), default="off", help="export --ocr に渡す値 (既定: off)")
    parser.add_argument("--timeout", type=int, default=DEFAULT_TIMEOUT_SEC, help="1 変換あたりのタイムアウト秒 (既定: 300)")
    parser.add_argument("--limit", type=int, default=None, help="処理するファイル数の上限")
    parser.add_argument("--baseline", help="比較対象の過去の report.json")
    parser.add_argument("--fail-on-regression", action="store_true", help="回帰検出時に終了コード 1 を返す (--baseline 必須)")
    parser.add_argument("--tolerance", type=float, default=DEFAULT_TOLERANCE_PCT_POINTS, help="回帰判定の許容 pt (既定: 1.0)")
    parser.add_argument("--jobs", type=int, default=1, help="並列変換数 (既定: 1)")
    parser.add_argument("--redact", action=argparse.BooleanOptionalAction, default=True,
                         help="report.json/.md 内のファイル名をハッシュ化する (既定: 有効)")
    return parser.parse_args(argv)


def main(argv: Optional[list] = None) -> int:
    args = parse_args(argv)

    corpus_root = Path(args.corpus).expanduser().resolve()
    if not corpus_root.is_dir():
        print(f"[corpus-regression] corpus directory not found: {corpus_root}", file=sys.stderr)
        return 2

    profiles = [item.strip() for item in args.profiles.split(",") if item.strip()]
    if not profiles or any(item not in ("readable", "roundtrip") for item in profiles):
        print("[corpus-regression] --profiles must be a comma-separated subset of readable,roundtrip", file=sys.stderr)
        return 2

    if args.fail_on_regression and not args.baseline:
        print("[corpus-regression] --fail-on-regression requires --baseline", file=sys.stderr)
        return 2
    if not math.isfinite(args.tolerance) or args.tolerance < 0:
        print("[corpus-regression] --tolerance must be finite and >= 0", file=sys.stderr)
        return 2
    if args.timeout <= 0:
        print("[corpus-regression] --timeout must be > 0", file=sys.stderr)
        return 2
    if args.limit is not None and args.limit <= 0:
        print("[corpus-regression] --limit must be > 0", file=sys.stderr)
        return 2
    if args.jobs < 1:
        print("[corpus-regression] --jobs must be >= 1", file=sys.stderr)
        return 2

    baseline_data = None
    if args.baseline:
        baseline_path = Path(args.baseline).expanduser()
        if not baseline_path.is_file():
            print(f"[corpus-regression] baseline report not found: {baseline_path}", file=sys.stderr)
            return 2
        try:
            baseline_data = json.loads(baseline_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            print(f"[corpus-regression] failed to read baseline: {exc}", file=sys.stderr)
            return 2
        if not isinstance(baseline_data, dict):
            print("[corpus-regression] baseline report must be a JSON object", file=sys.stderr)
            return 2
        if baseline_data.get("schema_version") != SCHEMA_VERSION:
            print(
                f"[corpus-regression] WARNING: baseline schema_version={baseline_data.get('schema_version')!r} "
                f"(expected {SCHEMA_VERSION!r}); comparison may be unreliable", file=sys.stderr,
            )

    try:
        cli_prefix, cli_kind = resolve_cli(args.cli)
    except CliResolutionError as exc:
        print(f"[corpus-regression] {exc}", file=sys.stderr)
        return 2
    cli_prefix, cli_kind = build_once(cli_prefix, cli_kind)

    probe = run_cli(cli_prefix + ["--version"], timeout=30)
    if probe["returncode"] != 0:
        print(f"[corpus-regression] CLI probe failed: {' '.join(cli_prefix)} --version", file=sys.stderr)
        print(f"[corpus-regression]   stderr: {probe['stderr'].strip()}", file=sys.stderr)
        return 2
    print(f"[corpus-regression] cli: {' '.join(cli_prefix)} ({cli_kind}); {probe['stdout'].strip()}")

    files = discover_corpus_files(corpus_root)
    if args.limit is not None:
        files = files[: args.limit]
    if not files:
        print(f"[corpus-regression] no .docx/.xlsx/.pptx/.pdf files found under {corpus_root}", file=sys.stderr)
        return 2
    print(f"[corpus-regression] {len(files)} file(s) discovered under {corpus_root}")

    out_dir = Path(args.out).expanduser().resolve() if args.out else (DEFAULT_OUT_ROOT / _timestamp())
    out_dir.mkdir(parents=True, exist_ok=True)

    mapping: dict = {}

    def process_one(path: Path) -> dict:
        return process_file(cli_prefix, corpus_root, path, out_dir, profiles, args.ocr, args.timeout, args.redact, mapping)

    results: list = [None] * len(files)  # type: ignore[list-item]
    if args.jobs > 1:
        with ThreadPoolExecutor(max_workers=args.jobs) as pool:
            futures = [pool.submit(process_one, path) for path in files]
            for index, future in enumerate(futures):
                results[index] = future.result()
                item = results[index]
                print(f"[corpus-regression] ({index + 1}/{len(files)}) {item['display_name']} "
                      f"format={item['format']} failed={item['failed']} warnings={len(item['warning_codes'])}")
    else:
        for index, path in enumerate(files):
            results[index] = process_one(path)
            item = results[index]
            print(f"[corpus-regression] ({index + 1}/{len(files)}) {item['display_name']} "
                  f"format={item['format']} failed={item['failed']} warnings={len(item['warning_codes'])}")

    aggregate = aggregate_results(results)

    report: dict = {
        "schema_version": SCHEMA_VERSION,
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "args": {
            "corpus": None if args.redact else str(corpus_root),
            "profiles": profiles,
            "ocr": args.ocr,
            "timeout": args.timeout,
            "limit": args.limit,
            "jobs": args.jobs,
            "redact": args.redact,
            "tolerance": args.tolerance,
        },
        "corpus": {
            "root": None if args.redact else str(corpus_root),
            "root_hash": hashlib.sha256(str(corpus_root).encode("utf-8")).hexdigest()[:16],
            "file_count": len(results),
        },
        "tooling": {"cli": cli_prefix, "cli_kind": cli_kind},
        "files": results,
        "aggregate": aggregate,
    }

    if baseline_data is not None:
        comparison = compute_baseline_comparison(report, baseline_data, args.tolerance)
        report["baseline_comparison"] = {
            "baseline_path": str(Path(args.baseline).expanduser().resolve()),
            "tolerance_pct_points": args.tolerance,
            **comparison,
        }

    if args.redact and mapping:
        (out_dir / "file_mapping.json").write_text(
            json.dumps(mapping, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8",
        )

    write_report_json(report, out_dir)
    write_report_md(report, out_dir)
    print_summary(report)
    print(f"[corpus-regression] report: {out_dir / 'report.json'} / {out_dir / 'report.md'}")

    comparison = report.get("baseline_comparison")
    if comparison and comparison["has_regression"]:
        print("[corpus-regression] REGRESSION DETECTED vs baseline (see report.md)", file=sys.stderr)
        if args.fail_on_regression:
            return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
